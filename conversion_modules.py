from PIL import Image,ImageDraw,ImageFont
import io
import pytesseract
from docx import Document
from fpdf import FPDF
import fitz
pytesseract.pytesseract.tesseract_cmd = r"C:\Program Files\Tesseract-OCR\tesseract.exe"


def ImageToPDF(file,ext):
    if ext in ["jpg","png","jpeg"]:
        image = Image.open(file.stream)
        if image.mode != "RGB":
            image = image.convert("RGB")
        pdf_bytes = io.BytesIO()
        image.save(pdf_bytes,format="PDF")
        pdf_bytes.seek(0)
        return pdf_bytes, "converted.pdf", "application/pdf"
    return None, None, None
    
def ImageToPNG(file,ext):
    if ext in ["jpg","jpeg"]:
        image = Image.open(file.stream)
        png_bytes = io.BytesIO()
        image.save(png_bytes, format="PNG")
        png_bytes.seek(0)
        return png_bytes, "converted.png", "image/png"
    return None, None, None

def ImageToJPG(file, ext):
    if ext == "png":
        image = Image.open(file.stream)
        if image.mode != "RGB":
            image = image.convert("RGB")
        jpg_bytes = io.BytesIO()
        image.save(
            jpg_bytes,
            format="JPEG",
            quality=95,
            optimize=True,
            subsampling=0
        )
        jpg_bytes.seek(0)
        return jpg_bytes, "converted.jpg", "image/jpeg"
    return None, None, None

def ImageToTXT(file,ext=0):
    image = Image.open(file.stream)
    image = image.convert("L")
    text = pytesseract.image_to_string(image)
    txt_bytes = io.BytesIO()
    txt_bytes.write(text.encode("utf-8"))
    txt_bytes.seek(0)
    return txt_bytes, "converted.txt", "text/plain"

def ImageToDOCX(file,ext=0):
    image = Image.open(file.stream)
    image = image.convert("L")
    text = pytesseract.image_to_string(image)
    doc = Document()
    doc.add_paragraph(text)
    doc_bytes = io.BytesIO()
    doc.save(doc_bytes)
    doc_bytes.seek(0)
    return doc_bytes, "converted.docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def txtToDocx(file):
    text = file.read().decode("utf-8")
    doc = Document()
    doc.add_heading("Converted File", 0)
    for line in text.split("\n"):
        doc.add_paragraph(line)
    output = io.BytesIO()
    doc.save(output)
    output.seek(0)
    return output, "converted.docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

def txtToPdf(file):
    text = file.read().decode("utf-8")
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Arial", size=12)
    for line in text.split("\n"):
        pdf.cell(0, 10, txt=line, ln=True)
    pdf_bytes = pdf.output(dest="S").encode("latin-1")
    output = io.BytesIO(pdf_bytes)
    output.seek(0)
    return output, "converted.pdf", "application/pdf"

def docxToTxt(file):
    doc = Document(file)
    text = ""
    for para in doc.paragraphs:
        text += para.text + "\n"
    output = io.BytesIO()
    output.write(text.encode("utf-8"))
    output.seek(0)
    return output, "converted.txt", "text/plain"

def pdfToTxt(file):
    pdf = fitz.open(stream=file.read(), filetype="pdf")
    text = ""
    for page in pdf:
        text += page.get_text()
    output = io.BytesIO()
    output.write(text.encode("utf-8"))
    output.seek(0)
    return output, "converted.txt", "text/plain"

def pdfTDocx(file):
    pdf = fitz.open(stream=file.read(), filetype="pdf")
    text = ""
    for page in pdf:
        text += page.get_text()
    doc = Document()
    doc.add_heading("Converted from PDF", 0)
    for line in text.split("\n"):
        doc.add_paragraph(line)
    output = io.BytesIO()
    doc.save(output)
    output.seek(0)
    return output, "converted.docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def pdfToImage(file,format):
    pdf = fitz.open(stream=file.read(), filetype="pdf")
    page = pdf[0]
    pix = page.get_pixmap()
    if format == "png":
        img_bytes = pix.tobytes("png")
        output = io.BytesIO(img_bytes)
        output.seek(0)
        return output, "converted.png", "image/png"
    img_bytes = pix.tobytes("jpeg")
    output = io.BytesIO(img_bytes)
    output.seek(0)
    return output, "converted.jpg", "image/jpeg"

def txtToImage(file,format):
    text = file.read().decode("utf-8")
    img = Image.new("RGB", (800, 1000), "white")
    draw = ImageDraw.Draw(img)
    font = ImageFont.load_default()
    y = 10
    for line in text.split("\n"):
        draw.text((10, y), line, fill="black", font=font)
        y += 20
    output = io.BytesIO()
    if format == "png":
        img.save(output, format="PNG")
        output.seek(0)
        return output, "converted.png", "image/png"
    img.save(output, format="JPEG")
    output.seek(0)
    return output, "converted.jpg", "image/jpeg"